"""Two runs of one version compared stage by stage: app.services.run_diff and its page."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pyarrow as pa
import pytest
from fastapi.testclient import TestClient

from app.core.files import save_upload
from app.core.frames import read_frame_table, write_frame_table
from app.main import app
from app.models.run_diff import RunComparison, StageComparison
from app.models.run_manifest import RunKind
from app.models.stage import stage_to_spec_dict
from app.runtime.manifest import resolve_output_path
from app.services import run as run_service, uploads, versioning
from app.services.errors import RunComparisonRefused
from app.services.run_diff import compare_output_tables, compare_runs
from app.services.workspace import resolve_run_dir
from app.tools.tutorial import TutorialAgentReference, TutorialContext, seed_tutorial_project

_TOUR_DATA = Path(__file__).resolve().parents[1] / "app" / "seeds" / "data"
_Q1 = _TOUR_DATA / "lda_data_Q1_2026.xlsx"
_INPUT_STAGE = "input_filings"
# test_seed_tutorial.py's cap: small, and the window still reaches the model stage.
_CAP = 50

# Every tour stage whose rows code computes and that a capped run finishes.
_COMPUTED_AND_FINISHED = {
    "input_filings", "find_ai_mentions", "keep_ai_candidates", "keep_ai_lobbying",
    "read_reported_money", "flag_in_house_filings", "ai_filings", "select_external_filings",
    "in_house_ai_filings", "in_house_ai_totals", "corpus_totals",
}
_JUDGED = {"judge_ai_substance", "review_ai_spend"}
# Downstream of the review queue, which halts a capped run.
_PENDING = {"confirm_ai_spend", "ai_spend_by_client", "ai_spend_totals"}

# 40 is the window's one AI-term row; no model reads income, so the seeded answers still apply.
_EDITED_ORDINALS = (0, 1, 2, 3, 40)
_EDITED_INCOME = "1234.00"
# The window holds no in-house filing, and corpus_totals counts filing ids alone.
_MISSED_BY_THE_EDIT = {"in_house_ai_filings", "in_house_ai_totals", "corpus_totals"}
_REACHED_BY_THE_EDIT = _COMPUTED_AND_FINISHED - _MISSED_BY_THE_EDIT


def test_runs_on_the_same_files_agree_and_a_changed_output_is_a_replay_violation(projects_root):
    tour = _seed_the_tour()
    first = _run_capped(tour, tour.input_files)
    second = _run_capped(tour, tour.input_files)

    agreeing = compare_runs(tour.project.id, first, second)
    _overwrite_first_cell(tour, second, "corpus_totals", "total_rows", _CAP + 1)
    violated = compare_runs(tour.project.id, first, second)

    assert agreeing.input_differences == [] and agreeing.find_replay_violations() == []
    assert _find_identical(agreeing) == _COMPUTED_AND_FINISHED
    assert _find_uncompared(agreeing) == _JUDGED | _PENDING
    assert [
        (stage.stage_id, [row.ordinal for row in stage.output_comparison.first_differing_rows])
        for stage in violated.find_replay_violations() if stage.output_comparison
    ] == [("corpus_totals", [0])]
    page = TestClient(app).get(f"/project/{tour.project.id}/runs/{first}/compare/{second}")
    assert "Both runs read the same files" in page.text
    assert 'href="#differs-corpus_totals">violation</a>' in page.text
    assert f'<td class="cell-mismatch" title="{_CAP + 1}">{_CAP + 1}</td>' in page.text
    runs_page = TestClient(app).get(f"/project/{tour.project.id}/runs")
    assert f"/runs/{second}/compare/{first}" in runs_page.text


def test_an_edited_input_file_is_named_and_what_it_reaches_differs_without_a_violation(
    projects_root, tmp_path
):
    tour = _seed_the_tour()
    before = _run_capped(tour, tour.input_files)
    edited_q1 = _upload(tour, _write_edited_q1(tmp_path))
    after = _run_capped(tour, {_INPUT_STAGE: [edited_q1, *tour.input_files[_INPUT_STAGE][1:]]})

    comparison = compare_runs(tour.project.id, before, after)

    assert [difference.stage_id for difference in comparison.input_differences] == [_INPUT_STAGE]
    assert _find_differing(comparison) == _REACHED_BY_THE_EDIT
    assert _find_identical(comparison) == _MISSED_BY_THE_EDIT
    assert comparison.find_replay_violations() == []
    loaded = _find_stage(comparison, _INPUT_STAGE).output_comparison
    assert loaded is not None and loaded.differing_row_count == len(_EDITED_ORDINALS)
    assert [
        (row.ordinal, row.differing_columns, (row.run_b_row or {})["income"])
        for row in loaded.first_differing_rows
    ] == [(ordinal, ["income"], _EDITED_INCOME) for ordinal in _EDITED_ORDINALS[:3]]
    page = TestClient(app).get(f"/project/{tour.project.id}/runs/{before}/compare/{after}")
    assert page.status_code == 200
    assert f'<span class="crumb-here" aria-current="page">Compared with {after}</span>' in page.text
    edited_hash = comparison.input_differences[0].run_b_read.files[0].sha256 or ""
    assert f"{_Q1.name} ({edited_hash[:8]})" in page.text
    assert all(
        f'href="#differs-{stage_id}">differs</a>' in page.text for stage_id in _REACHED_BY_THE_EDIT
    )
    assert f'<td class="cell-differs" title="{_EDITED_INCOME}">{_EDITED_INCOME}</td>' in page.text
    assert 'class="verdict verdict-fail"' not in page.text
    assert 'class="cell-mismatch"' not in page.text


def test_runs_pinned_to_different_versions_are_refused_naming_both(projects_root):
    tour = _seed_the_tour()
    first = _run_capped(tour, tour.input_files)
    restored = versioning.create_version_from_stages(
        tour.project.id,
        [stage_to_spec_dict(s) for s in versioning.load_version_stages(
            tour.project.id, tour.version_id)],
        message="the same stages, stored again",
    )
    second = _run_capped(tour, tour.input_files, version_id=restored.version_id)

    with pytest.raises(RunComparisonRefused) as refused:
        compare_runs(tour.project.id, first, second)

    assert tour.version_id in str(refused.value)
    assert restored.version_id in str(refused.value)
    page = TestClient(app).get(f"/project/{tour.project.id}/runs/{first}/compare/{second}")
    assert page.status_code == 409


def test_a_column_or_a_row_only_one_run_wrote_is_its_own_difference():
    first = pa.table({"name": ["a", "b"], "score": [1, 2]})
    second = pa.table({"name": ["a", "b", "c"], "score": [1, 3, 4], "note": ["x", "y", "z"]})

    compared = compare_output_tables(first, second)

    assert (compared.columns_only_in_run_a, compared.columns_only_in_run_b) == ([], ["note"])
    assert compared.differing_row_count == 2
    assert [
        (row.ordinal, row.run_a_row, row.differing_columns)
        for row in compared.first_differing_rows
    ] == [(1, {"name": "b", "score": 2}, ["score"]), (2, None, ["name", "score"])]


def _seed_the_tour() -> TutorialAgentReference:
    return seed_tutorial_project(TutorialContext(base_url="http://127.0.0.1:8788/"))


def _run_capped(
    tour: TutorialAgentReference, input_files: dict[str, list[str]],
    version_id: str | None = None,
) -> str:
    bindings = {
        stage_id: uploads.resolve_files_binding(tour.project.id, file_ids)
        for stage_id, file_ids in input_files.items()
    }
    manifest = run_service.execute(
        tour.project.id, version_id=version_id, bindings=bindings, limits={_INPUT_STAGE: _CAP})
    return str(manifest["run_id"])


def _write_edited_q1(directory: Path) -> Path:
    source = openpyxl.load_workbook(_Q1, read_only=True)
    edited = openpyxl.Workbook(write_only=True)
    sheet = edited.create_sheet()
    rows = source.worksheets[0].iter_rows(values_only=True)
    header = next(rows)
    sheet.append(header)
    income = header.index("income")
    for ordinal, row in enumerate(rows):
        values = list(row)
        if ordinal in _EDITED_ORDINALS:
            values[income] = _EDITED_INCOME
        sheet.append(values)
    path = directory / _Q1.name
    edited.save(path)
    return path


def _overwrite_first_cell(
    tour: TutorialAgentReference, run_id: str, stage_id: str, column: str, value: int
) -> None:
    manifest = run_service.read_run_manifest(tour.project.id, run_id, RunKind.production)
    record = manifest.find_stage_record(stage_id)
    assert record is not None
    path = resolve_output_path(
        resolve_run_dir(tour.project.id, run_id, RunKind.production), record.output_path)
    assert path is not None
    table = read_frame_table(path)
    index = table.column_names.index(column)
    cells = [value, *table.column(column).to_pylist()[1:]]
    write_frame_table(
        table.set_column(index, column, pa.array(cells, type=table.schema.field(column).type)),
        path)


def _upload(tour: TutorialAgentReference, path: Path) -> str:
    with path.open("rb") as handle:
        return save_upload(path.name, handle, tour.project.id).id


def _find_stage(comparison: RunComparison, stage_id: str) -> StageComparison:
    return next(stage for stage in comparison.stages if stage.stage_id == stage_id)


def _find_identical(comparison: RunComparison) -> set[str]:
    return {
        stage.stage_id for stage in comparison.stages
        if stage.output_comparison is not None and stage.output_comparison.is_identical
    }


def _find_differing(comparison: RunComparison) -> set[str]:
    return {
        stage.stage_id for stage in comparison.stages
        if stage.output_comparison is not None and not stage.output_comparison.is_identical
    }


def _find_uncompared(comparison: RunComparison) -> set[str]:
    return {stage.stage_id for stage in comparison.stages if stage.output_comparison is None}
